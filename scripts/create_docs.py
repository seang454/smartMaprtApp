import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_shading(cell, color_hex):
    """Set background color of a table cell (e.g. 'F1F5F9')."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set padding for a cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    """Set custom borders for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for border_name, border_props in borders.items():
        if border_props is not None:
            val, sz, color = border_props
            b_el = OxmlElement(f'w:{border_name}')
            b_el.set(qn('w:val'), val)
            b_el.set(qn('w:sz'), str(sz))
            b_el.set(qn('w:space'), '0')
            b_el.set(qn('w:color'), color)
            tcBorders.append(b_el)
        else:
            b_el = OxmlElement(f'w:{border_name}')
            b_el.set(qn('w:val'), 'none')
            tcBorders.append(b_el)
    tcPr.append(tcBorders)

def add_callout_box(doc, text_runs, bg_color="F8FAFC", border_color="CBD5E1", left_border_only=False, left_border_color="0284C7", left_border_sz=24):
    """Adds a callout box table with custom borders and background."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(6.8)
    set_cell_shading(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    if left_border_only:
        set_cell_borders(cell, 
                         top=('none', 0, 'auto'),
                         bottom=('none', 0, 'auto'),
                         left=('single', left_border_sz, left_border_color),
                         right=('none', 0, 'auto'))
    else:
        set_cell_borders(cell,
                         top=('single', 8, border_color),
                         bottom=('single', 8, border_color),
                         left=('single', 8, border_color),
                         right=('single', 8, border_color))
        
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    
    for item in text_runs:
        if isinstance(item, tuple):
            text = item[0]
            bold = item[1] if len(item) > 1 else False
            italic = item[2] if len(item) > 2 else False
            color = item[3] if len(item) > 3 else None
            size = item[4] if len(item) > 4 else Pt(9.5)
            font_name = item[5] if len(item) > 5 else "Calibri"
            
            run = p.add_run(text)
            run.bold = bold
            run.italic = italic
            run.font.name = font_name
            run.font.size = size
            if color:
                run.font.color.rgb = color
        elif isinstance(item, str):
            run = p.add_run(item)
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)
            
    # Add an empty spacer after the table
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(4)
    spacer.paragraph_format.line_spacing = 0.5
    return tbl

def add_code_block(doc, code_text):
    """Adds a formatted code snippet box."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(6.8)
    set_cell_shading(cell, "F8FAFC")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    set_cell_borders(cell,
                     top=('single', 6, 'E2E8F0'),
                     bottom=('single', 6, 'E2E8F0'),
                     left=('single', 6, 'E2E8F0'),
                     right=('single', 6, 'E2E8F0'))
    
    lines = code_text.strip().split('\n')
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(4)
    spacer.paragraph_format.line_spacing = 0.5

print('Helpers defined')
